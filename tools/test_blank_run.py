# -*- coding: utf-8 -*-
"""空白模板实跑验收（方案 2.5 节五条，源码级 GUI 实跑，不碰真实数据）。

  1. 空白模板启动：不崩、班级提示未设置、列表为空、状态栏提示导入
  2. 导入 5 名学生（2 室长、2 寝室）→ 模板重载 → 周表预览 5 行
  3. 生成第 1 周周表：标题、左右两栏（3+2）、签字栏追加、无残留学生
  4. 生成月表：M月总表 5 行、分数正确
  5. 寝室页 2 个寝室、可打分、应用后周表出现加减分

跑法：python tools/test_blank_run.py
"""

import importlib.util
import os
import shutil
import sys
import tempfile

import openpyxl

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BLANK = os.path.join(REPO_ROOT, "_build_pkg", "payload", "名单模板.xlsx")
sys.path.insert(0, os.path.join(REPO_ROOT, "软件说明"))

spec = importlib.util.spec_from_file_location(
    "moralapp", os.path.join(REPO_ROOT, "软件说明", "德育分自动化系统.py"))
mod = importlib.util.module_from_spec(spec)
sys.modules["moralapp"] = mod
spec.loader.exec_module(mod)

FAILS = []


def check(label, ok, extra=""):
    print(f"  [{'OK ' if ok else 'FAIL'}] {label}" + (f"  → {extra}" if extra else ""))
    if not ok:
        FAILS.append(label)


def pump(app, n=6):
    for _ in range(n):
        app.root.update_idletasks()
        app.root.update()


def main():
    workdir = tempfile.mkdtemp(prefix="blank-run-")
    shutil.copy2(BLANK, os.path.join(workdir, "名单模板.xlsx"))

    mod.MoralScoreApp._base_dir = lambda self: workdir   # 与 _ui-work 回归同款手法

    # 弹窗全部改成记录桩，避免测试被模态框挂住
    asked = []
    mod.messagebox.askyesno = lambda *a, **k: (asked.append(("askyesno", a)), False)[1]
    mod.messagebox.showinfo = lambda *a, **k: asked.append(("info", a))
    mod.messagebox.showwarning = lambda *a, **k: asked.append(("warn", a))
    mod.messagebox.showerror = lambda *a, **k: asked.append(("error", a))

    import tkinter as tk  # noqa: E402
    try:
        ctypes_dpi()
        root = tk.Tk()
        mod.enable_dpi_awareness()
        mod.apply_theme(root, accent=mod.ACCENT, accent_hover=mod.ACCENT_HOVER)
        mod.configure_app_styles(root)
        root.geometry("1240x820+30+30")
        app = mod.MoralScoreApp(root)
        app._update_busy = True          # 测试里不跑自动更新线程
        pump(app, 10)

        print("### 1. 空白模板启动")
        check("程序启动未崩", True)
        check("侧边栏班级提示未设置", "未设置" in app.lbl_class.cget("text"),
              app.lbl_class.cget("text"))
        check("周表预览为空", len(app.tree_week.get_children()) == 0)
        check("状态栏提示导入名单", "班级名单" in app.week_status.cget("text"),
              app.week_status.cget("text"))
        check("侧边栏版本号", app.lbl_version.cget("text") == f"v{mod.APP_VERSION}")
        check("月度预览为空", len(app.tree_month.get_children()) == 0)

        print("### 2. 导入 5 名学生（2 室长 / 2 寝室）")
        rows = [
            {"sid": "20250117001", "name": "张一", "dorm": "301", "role": "室长"},
            {"sid": "20250117002", "name": "李二", "dorm": "301", "role": "成员"},
            {"sid": "20250117003", "name": "王三", "dorm": "302", "role": "成员"},
            {"sid": "20250117004", "name": "赵四", "dorm": "302", "role": "室长"},
            {"sid": "20250117005", "name": "钱五", "dorm": "", "role": "成员"},
        ]
        app._apply_roster_to_template(rows, "测试技术999")
        pump(app, 10)
        check("导入后预览 5 行", len(app.tree_week.get_children()) == 5,
              str(len(app.tree_week.get_children())))
        check("班级名生效", app.class_name == "测试技术999班", app.class_name)
        check("侧边栏班级更新", app.lbl_class.cget("text") == "测试技术999班")
        check("左右布局 3+2", len(app.base_layout) == 3 and app.base_layout[2][1] == "",
              str(app.base_layout))
        check("寝室视图 2 个寝室", len(app.dorm_stars) == 2, str(list(app.dorm_stars)))
        check("模板备份存在", any(n.startswith("名单模板.自动备份-")
                                  for n in os.listdir(workdir)))

        print("### 3. 生成第 1 周周表")
        app.current_week.set("1")
        app.save_week_data()
        pump(app, 8)
        folder = os.path.join(workdir, "2026年10月")
        wfile = os.path.join(folder, app.get_week_file_name())
        check("周表文件生成", os.path.exists(wfile), wfile)
        wb = openpyxl.load_workbook(wfile)
        ws = wb["第1周"]
        check("标题班级替换", str(ws["A1"].value) == "测试技术999班10月1周记实",
              str(ws["A1"].value))
        l_ids = [ws.cell(4 + k, 1).value for k in range(3)]
        r_ids = [ws.cell(4 + k, 5).value for k in range(3)]
        check("左栏 3 人", l_ids == ["20250117001", "20250117002", "20250117003"],
              str(l_ids))
        check("右栏 2 人（第 3 对右空）", r_ids[:2] == ["20250117004", "20250117005"]
              and r_ids[2] is None, str(r_ids))
        check("签字栏已追加", any("班主任签名" == str(ws.cell(r, 1).value)
                              for r in range(5, ws.max_row + 1)))
        wb.close()

        print("### 5. 寝室打分 → 应用到周表")
        app._set_dorm_star("301", 5)     # 五星：周 成员+1
        app._set_dorm_star("302", 3)     # 三星：周 成员-1 室长-1.5
        app.apply_dorm_scores()
        pump(app, 8)
        app.save_week_data()
        pump(app, 6)
        wb = openpyxl.load_workbook(wfile)
        ws = wb["第1周"]
        # 张一(301室长,五星) 李二(301成员,五星) 王三(302成员,三星) 赵四(302室长,三星)
        row_score = {}
        for r in range(4, 10):
            for sid_col, sc_col, rs_col in ((1, 3, 4), (5, 7, 8)):   # 左栏 / 右栏
                sid = ws.cell(r, sid_col).value
                if sid:
                    row_score[str(sid)] = (ws.cell(r, sc_col).value,
                                           str(ws.cell(r, rs_col).value))
        check("301 五星生效", row_score.get("20250117001", (None,))[0] == 1,
              str(row_score.get("20250117001")))
        check("302 三星成员 -1", row_score.get("20250117003", (None,))[0] == -1,
              str(row_score.get("20250117003")))
        check("302 三星室长 -1.5", row_score.get("20250117004", (None,))[0] == -1.5,
              str(row_score.get("20250117004")))
        wb.close()

        print("### 4. 生成月表（首月模式，底分 80）")
        app.is_first_month.set(True)
        app.month_sum = {r["sid"]: (1.0 if r["dorm"] == "301" else
                                    (-1.0 if r["role"] == "成员" else -1.5))
                         for r in rows}
        app.week_scanned = True
        app.last_month_loaded = True
        app.save_month_data()
        pump(app, 8)
        mfile = os.path.join(folder, app.get_month_file_name())
        check("月表文件生成", os.path.exists(mfile), mfile)
        wb = openpyxl.load_workbook(mfile)
        check("sheet 名 10月总表", "10月总表" in wb.sheetnames, str(wb.sheetnames))
        ws = wb["10月总表"]
        got = {}
        for r in range(2, 10):
            sid = ws.cell(r, 1).value
            if sid:
                got[str(sid)] = ws.cell(r, 5).value
        check("月表 5 行", len(got) == 5, str(len(got)))
        check("张一 80+1=81", got.get("20250117001") == 81, str(got.get("20250117001")))
        check("王三 80-1=79", got.get("20250117003") == 79, str(got.get("20250117003")))
        check("赵四 80-1.5=78.5", got.get("20250117004") == 78.5,
              str(got.get("20250117004")))
        check("D 列月份写入", ws.cell(2, 4).value == "2026年10月",
              str(ws.cell(2, 4).value))
        wb.close()

        # 用户数据隔离：真实安装目录没有被碰
        check("工作区与真实目录隔离", workdir.startswith(tempfile.gettempdir()))

        root.destroy()
    finally:
        shutil.rmtree(workdir, ignore_errors=True)

    if FAILS:
        print(f"\n[失败] {len(FAILS)} 项：{FAILS}")
        return 1
    print("\n[通过] 空白模板五条验收全部通过。")
    return 0


def ctypes_dpi():
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass


if __name__ == "__main__":
    sys.exit(main())
