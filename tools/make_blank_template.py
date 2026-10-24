# -*- coding: utf-8 -*-
"""生成空白版《名单模板.xlsx》（分发用）。

铁律：本脚本对源模板【只读】，只写到打包暂存区，绝不 save 回源路径。
源文件是用户正在用的真实数据（学生姓名 + 分数结构），改坏就是事故。

用法：
    python tools/make_blank_template.py [源模板.xlsx] [输出.xlsx]
不带参数时用下面的默认路径。

抹除规则见《方案-安装包与热更新.md》第 2 节 / 9.1 节：
  - 周表模板   A1 去班级名（保留 #月#周 占位），第 4 行起全部值清空
  - 月表模板   第 2 行起清 A/B/C/E，D 列 #/#/# 占位保留
  - 寝室名单   第 3 行起（第一个寝室行）全部值清空
  - 周表模板2  签字值列（D/H，真实人名）清空，标签保留
  - 寝室加分规则 / 加分预览  原样保留（业务规则）
结构（合并单元格、列宽、行高、样式）全部保留。
"""

import os
import sys

import openpyxl

# 源模板：用户真实在用的文件（只读！）
DEFAULT_SRC = r"E:\邮电\权益部分\德育分统计表\德育分自动化系统\名单模板.xlsx"

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_DST = os.path.join(REPO_ROOT, "_build_pkg", "payload", "名单模板.xlsx")


def clear_area(ws, row_start, cols):
    """把从 row_start 起的指定列的值清空（跳过 MergedCell，保留样式与骨架）。"""
    for r in range(row_start, ws.max_row + 1):
        for c in cols:
            cell = ws.cell(row=r, column=c)
            if type(cell).__name__ == "MergedCell":
                continue
            cell.value = None


def make_blank(src, dst):
    wb = openpyxl.load_workbook(src)   # 默认保留样式、合并、列宽

    # 1) 周表模板：A1 去班级名；第 4 行起所有列清空（含残留加减分/原因）
    ws = wb["周表模板"]
    ws["A1"] = "#月#周记实"
    clear_area(ws, 4, range(1, 9))          # A..H

    # 2) 月表模板：第 2 行起清 A(学号) B(姓名) C(班级) E(积分)；D 的 #/#/# 保留
    ws = wb["月表模板"]
    clear_area(ws, 2, (1, 2, 3, 5))

    # 3) 寝室名单：表头占 1~2 行，第 3 行起是第一个寝室 → 全部清空
    ws = wb["寝室名单"]
    clear_area(ws, 3, range(1, 8))          # A..G

    # 4) 周表模板2：签字栏 D/H 列是真实签名（考评小组 5 人、班主任、辅导员），
    #    分发版清掉这些值，只留「考评小组签名/班主任签名」等标签
    ws = wb["周表模板2"]
    clear_area(ws, 1, (4, 8))               # D、H 列

    # 5) 寝室加分规则 / 加分预览：原样不动

    os.makedirs(os.path.dirname(dst), exist_ok=True)
    wb.save(dst)
    wb.close()
    return dst


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SRC
    dst = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_DST
    if not os.path.exists(src):
        print(f"[错误] 找不到源模板：{src}")
        return 1
    make_blank(src, dst)
    print(f"[OK] 空白模板已生成：{dst}")
    print(f"     源模板（未改动）：{src}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
